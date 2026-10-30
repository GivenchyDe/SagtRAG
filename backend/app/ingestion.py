"""Phase 2 文档格式验证、文本提取与稳定分块。"""

import hashlib
import io
import zipfile
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}
PARSER_VERSION = "sagtrag-parser-1"
CHUNKER_VERSION = "paragraph-window-1"


class DocumentInputError(ValueError):
    """表示上传文件格式、签名或内容不可处理。"""


def validate_document_bytes(filename: str, media_type: str, content: bytes, max_bytes: int) -> str:
    """校验扩展名、允许 MIME、大小和文件容器特征并返回规范化格式。"""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise DocumentInputError("不支持此文件格式")
    if not content or len(content) > max_bytes:
        raise DocumentInputError("文件为空或超过上传大小限制")
    expected = {
        ".pdf": {"application/pdf", "application/octet-stream"},
        ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
        ".txt": {"text/plain", "application/octet-stream"},
        ".md": {"text/markdown", "text/plain", "application/octet-stream"},
        ".markdown": {"text/markdown", "text/plain", "application/octet-stream"},
    }[suffix]
    if media_type.lower().split(";", 1)[0].strip() not in expected:
        raise DocumentInputError("文件类型与扩展名不匹配")
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise DocumentInputError("PDF 文件签名无效")
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = archive.namelist()
                total_size = sum(item.file_size for item in archive.infolist())
                if len(names) > 500 or total_size > max_bytes * 40 or "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise DocumentInputError("DOCX 容器结构无效或超出处理限制")
        except (zipfile.BadZipFile, OSError) as exc:
            raise DocumentInputError("DOCX 文件容器损坏") from exc
    if suffix in {".txt", ".md", ".markdown"}:
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentInputError("文本文件必须使用 UTF-8 编码") from exc
    return suffix


def extract_paragraphs(filename: str, content: bytes) -> list[str]:
    """从 PDF、DOCX 或 UTF-8 文本提取非空段落；扫描 PDF 不执行 OCR。"""
    suffix = PurePath(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise DocumentInputError("加密 PDF 暂不支持")
            paragraphs = [part.strip() for page in reader.pages for part in (page.extract_text() or "").splitlines() if part.strip()]
            if not paragraphs:
                raise DocumentInputError("PDF 未提取到可搜索文本，扫描件暂不支持 OCR")
            return paragraphs
        if suffix == ".docx":
            document = DocxDocument(io.BytesIO(content))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    value = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if value:
                        paragraphs.append(value)
            if not paragraphs:
                raise DocumentInputError("DOCX 文件中没有可提取文本")
            return paragraphs
        text = content.decode("utf-8-sig")
        paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
        if not paragraphs:
            raise DocumentInputError("文本文件中没有内容")
        return paragraphs
    except DocumentInputError:
        raise
    except Exception as exc:
        raise DocumentInputError("文件无法解析或结构损坏") from exc


def split_into_chunks(paragraphs: list[str], *, revision_id: str, size: int = 1200, overlap: int = 150) -> list[dict[str, object]]:
    """优先按段落边界合并，超长段落再按字符窗口重叠切分。"""
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > size:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(paragraph):
                end = min(start + size, len(paragraph))
                chunks.append(paragraph[start:end])
                if end == len(paragraph):
                    break
                start = end - overlap
            continue
        candidate = f"{current}\n\n{paragraph}" if current else paragraph
        if len(candidate) > size and current:
            chunks.append(current)
            current = paragraph
        else:
            current = candidate
    if current:
        chunks.append(current)
    return [
        {
            "ordinal": ordinal,
            "content": value,
            "chunk_key": hashlib.sha256(f"{revision_id}:{ordinal}:{value}".encode("utf-8")).hexdigest(),
        }
        for ordinal, value in enumerate(chunks)
    ]
